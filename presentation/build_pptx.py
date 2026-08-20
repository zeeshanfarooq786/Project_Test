#!/usr/bin/env python3
"""18-slide Ring Green Pakistan deck: Mülltrennung & Pfand."""

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

W, H = Inches(13.333), Inches(7.5)
ASSETS = Path(__file__).resolve().parent / "assets"

GREEN = RGBColor(0x2E, 0x7D, 0x32)
NAVY = RGBColor(0x0D, 0x47, 0xA1)
BG = RGBColor(0xF8, 0xF9, 0xFA)
YELLOW = RGBColor(0xFB, 0xC0, 0x2D)
INK = RGBColor(0x21, 0x21, 0x21)
MUTED = RGBColor(0x5F, 0x6B, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DANGER = RGBColor(0xC6, 0x28, 0x28)
BLUEB = RGBColor(0x15, 0x65, 0xC0)
BROWN = RGBColor(0x6D, 0x4C, 0x41)
GRAY = RGBColor(0x45, 0x5A, 0x64)
CREAM = RGBColor(0xFF, 0xF8, 0xE1)


def set_run(run, size, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", font)


def txt(shape, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return tf


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill(s, color)
    return s


def rrect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(s, color)
    try:
        s.adjustments[0] = 0.08
    except Exception:
        pass
    return s


def box(slide, l, t, w, h, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT):
    b = slide.shapes.add_textbox(l, t, w, h)
    txt(b, text, size, bold, color, align)
    return b


def bullets(slide, items, l, t, w, h, size=14, color=INK):
    b = slide.shapes.add_textbox(l, t, w, h)
    tf = b.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size, False, color)
    return b


def footer(slide, n):
    box(slide, Inches(0.4), Inches(7.18), Inches(8), Inches(0.22), "RING GREEN PAKISTAN  ·  4TH SESSION", 10, True, MUTED)
    box(slide, Inches(10.4), Inches(7.18), Inches(2.5), Inches(0.22), f"{n:02d}  /  18", 10, True, MUTED, PP_ALIGN.RIGHT)


def blank():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, W, H, BG)
    rect(slide, 0, 0, W, Inches(0.08), GREEN)
    rect(slide, 0, Inches(7.42), W, Inches(0.08), YELLOW)
    return slide


def img(slide, name, l, t, w, h):
    p = ASSETS / name
    if p.exists():
        slide.shapes.add_picture(str(p), l, t, w, h)


def card(slide, l, t, w, h):
    return rrect(slide, l, t, w, h, WHITE)


def bin_slide(n, title, sub, allow, forbid, note, image, accent):
    s = blank()
    rect(s, 0, 0, Inches(0.18), H, accent)
    box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), title.upper(), 12, True, accent)
    box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.5), title, 26, True, INK)
    box(s, Inches(0.45), Inches(1.0), Inches(12), Inches(0.35), sub, 14, False, MUTED)
    card(s, Inches(0.45), Inches(1.5), Inches(6.0), Inches(3.55))
    box(s, Inches(0.65), Inches(1.65), Inches(5.6), Inches(0.3), "ALLOWED", 12, True, GREEN)
    bullets(s, allow, Inches(0.65), Inches(2.05), Inches(5.6), Inches(2.8), 15)
    card(s, Inches(6.65), Inches(1.5), Inches(6.2), Inches(3.55))
    box(s, Inches(6.85), Inches(1.65), Inches(5.8), Inches(0.3), "FORBIDDEN", 12, True, DANGER)
    bullets(s, forbid, Inches(6.85), Inches(2.05), Inches(5.8), Inches(2.8), 15, DANGER)
    rrect(s, Inches(0.45), Inches(5.2), Inches(8.3), Inches(1.7), CREAM)
    box(s, Inches(0.65), Inches(5.4), Inches(7.9), Inches(1.35), note, 14, False, INK)
    img(s, image, Inches(8.95), Inches(5.2), Inches(3.9), Inches(1.7))
    footer(s, n)


prs = Presentation()
prs.slide_width = W
prs.slide_height = H

# 1 Title
s = prs.slides.add_slide(prs.slide_layouts[6])
img(s, "title-hero.jpg", 0, 0, W, H)
rect(s, 0, 0, Inches(8.2), H, RGBColor(0x0D, 0x1B, 0x12))
box(s, Inches(0.5), Inches(0.4), Inches(7), Inches(0.3), "RING GREEN PAKISTAN  ·  4TH SESSION", 12, True, YELLOW)
box(s, Inches(0.5), Inches(1.5), Inches(7.2), Inches(1.6), "German Waste Management System", 36, True, WHITE)
box(s, Inches(0.5), Inches(3.2), Inches(7), Inches(0.4), "Mülltrennung  &  Pfand", 18, True, YELLOW)
box(s, Inches(0.5), Inches(3.7), Inches(7), Inches(1.1), "Welcome. Learn how Germany sorts waste as a resource — and how Pakistan can adapt it at home, in streets, and in law.", 16, False, WHITE)
for i, (t, d) in enumerate([
    ("OBJECTIVE 01", "Understand every German bin + Pfand loop"),
    ("OBJECTIVE 02", "Map a 4-phase Pakistan adaptation roadmap"),
    ("SPEAKER", "Ring Green Pakistan · Session Facilitators"),
]):
    x = Inches(0.5) + i * Inches(4.2)
    rrect(s, x, Inches(5.2), Inches(4.0), Inches(1.4), RGBColor(0x1B, 0x5E, 0x20))
    box(s, x + Inches(0.15), Inches(5.3), Inches(3.7), Inches(0.3), t, 11, True, YELLOW)
    box(s, x + Inches(0.15), Inches(5.65), Inches(3.7), Inches(0.75), d, 14, True, WHITE)
footer(s, 1)

# 2 Intro
s = blank()
box(s, Inches(0.45), Inches(0.25), Inches(12), Inches(0.3), "MÜLLTRENNUNG — WHY GERMANY LEADS", 12, True, GREEN)
box(s, Inches(0.45), Inches(0.55), Inches(12), Inches(0.55), "Germany treats waste as a resource — not leftover dirt", 26, True, INK)
card(s, Inches(0.45), Inches(1.3), Inches(3.9), Inches(1.7))
box(s, Inches(0.6), Inches(1.4), Inches(3.6), Inches(0.7), "69%+", 32, True, GREEN)
box(s, Inches(0.6), Inches(2.1), Inches(3.6), Inches(0.7), "Municipal waste recycled. Europe’s leader by household discipline.", 13, False, MUTED)
card(s, Inches(4.5), Inches(1.3), Inches(3.9), Inches(1.7))
box(s, Inches(4.65), Inches(1.4), Inches(3.6), Inches(0.7), "98%", 32, True, GREEN)
box(s, Inches(4.65), Inches(2.1), Inches(3.6), Inches(0.7), "Pfand bottle return. Deposit makes returning natural.", 13, False, MUTED)
card(s, Inches(0.45), Inches(3.15), Inches(7.95), Inches(3.7))
box(s, Inches(0.65), Inches(3.3), Inches(7.5), Inches(0.3), "THE CORE PHILOSOPHY", 12, True, NAVY)
bullets(s, [
    "Waste is a raw material sitting in the wrong place.",
    "Sorting is mandatory at home — the kitchen is the first factory.",
    "Colour is the UI: Blue / Yellow / Brown / Gray + Glass + Pfand.",
    "The 4-bin RNO standard keeps streams uncontaminated.",
], Inches(0.65), Inches(3.7), Inches(7.5), Inches(2.9), 16)
img(s, "german-bins-street.jpg", Inches(8.6), Inches(1.3), Inches(4.3), Inches(5.55))
footer(s, 2)

bin_slide(3, "Blue Bin — Paper & Cardboard", "Clean fibre only. Grease and food kill the recycling loop.",
          ["Clean paper & newspapers", "Magazines & catalogues", "Flattened cardboard boxes", "Envelopes (windows OK in most cities)"],
          ["Dirty pizza boxes (grease / cheese)", "Used tissues & kitchen paper", "Waxed / coated paper", "Wet or food-stained fibre"],
          "Key rule: flatten every cardboard box. Volume, not weight, fills the communal bin.",
          "blue-bin-paper.jpg", BLUEB)

bin_slide(4, "Yellow Bin / Yellow Bag — Packaging", "Plastic, metal, composite. Spoon-clean — do not waste water rinsing.",
          ["Yogurt cups (plastic packaging)", "Tin cans (metal)", "Aluminium foil if mostly clean", "TetraPak / milk cartons"],
          ["Food-filled packs", "Pfand bottles/cans with DPG logo (return to supermarket)", "Dirty diapers or residual waste", "Do not waste litres of water rinsing empty packs"],
          "Spoon-clean rule: scrape, don’t rinse litres of water. DPG-logo bottles go to the Pfandautomat, not the yellow bag.",
          "yellow-bin-packaging.jpg", YELLOW)

bin_slide(5, "Brown Bin — Organic / Biotonne", "Kitchen and garden leftovers become compost and biogas.",
          ["Fruit & vegetable peels", "Food scraps", "Coffee grounds & tea bags", "Garden trimmings"],
          ["Plastic bags — even “compostable” bio-plastic", "Use newspaper or paper bags only", "Glass, metal, packaging", "Residual sanitary waste"],
          "Environmental value: kitchen → Biotonne → municipal composting + biogas power. Organic waste is fuel and fertiliser.",
          "brown-bin-organic.jpg", BROWN)

# 6 Gray
s = blank()
rect(s, 0, 0, Inches(0.18), H, GRAY)
box(s, Inches(0.45), Inches(0.25), Inches(12), Inches(0.3), "GRAY BIN · RESTMÜLL", 12, True, GRAY)
box(s, Inches(0.45), Inches(0.55), Inches(12), Inches(0.55), "Black / Gray Bin — residual waste only", 26, True, INK)
card(s, Inches(0.45), Inches(1.3), Inches(6.3), Inches(5.5))
box(s, Inches(0.65), Inches(1.45), Inches(5.9), Inches(0.3), "ALLOWED (NON-RECYCLABLES)", 12, True, NAVY)
bullets(s, ["Vacuum bags", "Ash & cigarette butts", "Diapers", "Broken ceramics", "Sanitation items", "Dirty mix that cannot be recycled"], Inches(0.65), Inches(1.9), Inches(5.9), Inches(3.2), 16)
box(s, Inches(0.65), Inches(5.3), Inches(5.9), Inches(1.2), "Golden rule: keep this bin as empty as possible by using the other streams first.", 15, True, INK)
rrect(s, Inches(7.0), Inches(1.3), Inches(5.85), Inches(5.5), RGBColor(0x26, 0x32, 0x38))
box(s, Inches(7.2), Inches(1.45), Inches(5.45), Inches(0.3), "FINAL DESTINATION", 12, True, YELLOW)
box(s, Inches(7.2), Inches(1.85), Inches(5.45), Inches(0.5), "Incineration → district heating", 20, True, WHITE)
box(s, Inches(7.2), Inches(2.45), Inches(5.45), Inches(1.5), "Restmüll is burned in controlled plants. Heat warms homes. A full gray bin is treated as a sorting failure, not a normal week.", 15, False, WHITE)
img(s, "gray-bin-residual.jpg", Inches(7.2), Inches(4.15), Inches(5.45), Inches(2.4))
footer(s, 6)

# 7 Glass
s = blank()
box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "GLASCONTAINER", 12, True, GREEN)
box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.5), "Street glass igloos — sort by colour, not by brand", 24, True, INK)
cols = [
    (RGBColor(0xEC, 0xEF, 0xF1), INK, "White · Weißglas", "Clear bottles and jars. Colour contamination ruins the melt."),
    (GREEN, WHITE, "Green · Grünglas", "Green plus blue and other mixed colours in most cities."),
    (BROWN, WHITE, "Brown · Braunglas", "Brown beer and ordinary household brown bottles."),
]
for i, (bgc, fg, title, body) in enumerate(cols):
    x = Inches(0.45) + i * Inches(4.2)
    rrect(s, x, Inches(1.2), Inches(4.0), Inches(2.0), bgc)
    box(s, x + Inches(0.2), Inches(1.35), Inches(3.6), Inches(0.4), title, 16, True, fg)
    box(s, x + Inches(0.2), Inches(1.8), Inches(3.6), Inches(1.15), body, 13, False, fg)
card(s, Inches(0.45), Inches(3.4), Inches(7.3), Inches(3.4))
box(s, Inches(0.65), Inches(3.55), Inches(6.9), Inches(0.3), "HOUSEHOLD VS SPECIAL GLASS", 12, True, NAVY)
bullets(s, [
    "Jam jars, sauce bottles, ordinary drink glass → igloo",
    "Medical / pharma vials → Restmüll or hazardous stream",
    "Pyrex / heat-resistant cookware → Restmüll (different melting point)",
    "Quiet hours: weekdays daytime only. No Sundays, no nights.",
], Inches(0.65), Inches(4.0), Inches(6.9), Inches(2.5), 15)
img(s, "glass-igloos.jpg", Inches(8.0), Inches(3.4), Inches(4.85), Inches(3.4))
footer(s, 7)

# 8 Pfand
s = blank()
box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "PFAND · DEPOSIT-REFUND", 12, True, GREEN)
box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.5), "The Pfand loop — 98% of bottles come home", 24, True, INK)
steps = ["1. Customer pays deposit", "2. Uses the drink", "3. Pfandautomat at supermarket", "4. Cash / voucher back"]
for i, t in enumerate(steps):
    x = Inches(0.45) + i * Inches(3.2)
    rrect(s, x, Inches(1.15), Inches(3.0), Inches(1.15), NAVY if i % 2 == 0 else GREEN)
    box(s, x + Inches(0.12), Inches(1.4), Inches(2.76), Inches(0.7), t, 14, True, WHITE, PP_ALIGN.CENTER)
rrect(s, Inches(0.45), Inches(2.55), Inches(6.3), Inches(2.0), NAVY)
box(s, Inches(0.65), Inches(2.7), Inches(5.9), Inches(0.3), "SINGLE-USE  ·  EINWEG", 12, True, YELLOW)
box(s, Inches(0.65), Inches(3.05), Inches(5.9), Inches(0.55), "€0.25", 28, True, WHITE)
box(s, Inches(0.65), Inches(3.65), Inches(5.9), Inches(0.7), "PET bottles & cans with the DPG logo.", 14, False, WHITE)
rrect(s, Inches(0.45), Inches(4.7), Inches(6.3), Inches(2.05), GREEN)
box(s, Inches(0.65), Inches(4.85), Inches(5.9), Inches(0.3), "MULTI-USE  ·  MEHRWEG", 12, True, YELLOW)
box(s, Inches(0.65), Inches(5.2), Inches(5.9), Inches(0.5), "€0.08 – €0.15", 26, True, WHITE)
box(s, Inches(0.65), Inches(5.75), Inches(5.9), Inches(0.8), "Heavy plastic or glass beer/water bottles, reused up to 50 times.", 14, False, WHITE)
img(s, "pfandautomat.jpg", Inches(7.0), Inches(2.55), Inches(5.85), Inches(4.2))
footer(s, 8)

# 9 Apartments
s = blank()
box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "APARTMENT BUILDINGS", 12, True, GREEN)
box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.55), "One neighbour’s lazy sorting becomes everyone’s bill", 24, True, INK)
img(s, "apartment-bins.jpg", Inches(0.45), Inches(1.25), Inches(6.1), Inches(5.55))
card(s, Inches(6.75), Inches(1.25), Inches(6.1), Inches(1.55))
box(s, Inches(6.95), Inches(1.35), Inches(5.7), Inches(0.3), "HOW IT WORKS", 12, True, NAVY)
box(s, Inches(6.95), Inches(1.7), Inches(5.7), Inches(0.9), "Multi-family houses share a central bin yard. Collection is a building service, not a private hobby.", 14, False, INK)
rrect(s, Inches(6.75), Inches(2.95), Inches(6.1), Inches(2.0), RGBColor(0xFF, 0xEB, 0xEE))
box(s, Inches(6.95), Inches(3.05), Inches(5.7), Inches(0.3), "SHARED RESPONSIBILITY", 12, True, DANGER)
box(s, Inches(6.95), Inches(3.4), Inches(5.7), Inches(1.4), "If 1 tenant mis-sorts, the municipality can charge higher non-sorted fees to the whole building. Landlord spreads cost via Nebenkosten — or warns the renter.", 14, False, INK)
rrect(s, Inches(6.75), Inches(5.1), Inches(6.1), Inches(1.7), NAVY)
box(s, Inches(6.95), Inches(5.25), Inches(5.7), Inches(1.35), "Repeat violations damage the residential record, risk lease termination, or trigger municipal citations.", 14, False, WHITE)
footer(s, 9)

# 10 Laws DE
s = blank()
box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "MONITORING · LAWS · FINES", 12, True, GREEN)
box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.5), "The state watches the bin — and will leave it standing", 24, True, INK)
card(s, Inches(0.45), Inches(1.2), Inches(6.3), Inches(2.3))
box(s, Inches(0.65), Inches(1.35), Inches(5.9), Inches(0.3), "GOVERNMENT MECHANISM", 12, True, NAVY)
bullets(s, ["Waste inspectors — Mülldetektive", "Abfallkalender pickup calendar by street", "Separate trucks per stream"], Inches(0.65), Inches(1.75), Inches(5.9), Inches(1.55), 15)
card(s, Inches(0.45), Inches(3.65), Inches(6.3), Inches(3.1))
box(s, Inches(0.65), Inches(3.8), Inches(5.9), Inches(0.3), "COLLECTION RULES", 12, True, NAVY)
box(s, Inches(0.65), Inches(4.2), Inches(5.9), Inches(1.2), "Wrong materials → warning sticker and no collection until re-sorted.", 15, False, INK)
box(s, Inches(0.65), Inches(5.4), Inches(5.9), Inches(1.1), "Illegal dumping (Sperrmüll / Wildmüll): fines €50 up to €2,500+ by federal state.", 15, True, DANGER)
img(s, "waste-truck.jpg", Inches(7.0), Inches(1.2), Inches(5.85), Inches(5.55))
footer(s, 10)

# 11 Journey
s = blank()
box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "WASTE JOURNEY", 12, True, GREEN)
box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.5), "Five steps from kitchen to circular re-entry", 24, True, INK)
js = [
    ("1", "Source separation", "Kitchen caddies: paper, packaging, organic, residual."),
    ("2", "Bin placement", "Curbside on the Abfallkalender day only."),
    ("3", "Specialized pickup", "Dedicated trucks per stream. No mixed lorry."),
    ("4", "Sorting & Pfand", "Optical plants + reverse vending machines."),
    ("5", "Circular re-entry", "Compost, pellets, glass remelt, bottle refill."),
]
for i, (n, title, body) in enumerate(js):
    x = Inches(0.35) + i * Inches(2.58)
    rrect(s, x, Inches(1.3), Inches(2.45), Inches(4.5), WHITE)
    rect(s, x, Inches(1.3), Inches(2.45), Inches(0.1), GREEN if i != 2 else YELLOW)
    box(s, x + Inches(0.15), Inches(1.55), Inches(2.15), Inches(0.45), n, 22, True, NAVY, PP_ALIGN.CENTER)
    box(s, x + Inches(0.12), Inches(2.15), Inches(2.2), Inches(1.0), title, 16, True, INK, PP_ALIGN.CENTER)
    box(s, x + Inches(0.12), Inches(3.2), Inches(2.2), Inches(2.2), body, 13, False, MUTED, PP_ALIGN.CENTER)
rrect(s, Inches(0.45), Inches(6.0), Inches(12.4), Inches(0.95), CREAM)
box(s, Inches(0.65), Inches(6.2), Inches(12.0), Inches(0.6), "If step 1 fails, steps 3–5 become expensive theatre. Pakistan’s work begins in the kitchen.", 15, True, INK)
footer(s, 11)

# 12 Pakistan reality
s = blank()
box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "PAKISTAN · CURRENT REALITY", 12, True, DANGER)
box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.55), "We already collect — mixed, late, and often burning", 24, True, INK)
challenges = [
    ("Mixed dumping", "Everything in one bag. Plastic/glass value dies on contact with wet waste."),
    ("Landfills & nullahs", "Unorganized dumps and drain blockages. Monsoon turns trash into floods."),
    ("Open burning", "Neighbourhood piles become toxic smoke — a hospital invoice."),
    ("Informal risk", "Kabbadiwalas work without PPE in contaminated heaps."),
]
for (t, b), (x, y) in zip(challenges, [(0.45, 1.25), (0.45, 3.95), (3.65, 1.25), (3.65, 3.95)]):
    rrect(s, Inches(x), Inches(y), Inches(3.05), Inches(2.5), RGBColor(0xFF, 0xEB, 0xEE))
    box(s, Inches(x + 0.15), Inches(y + 0.15), Inches(2.75), Inches(0.45), t, 16, True, DANGER)
    box(s, Inches(x + 0.15), Inches(y + 0.7), Inches(2.75), Inches(1.55), b, 13, False, INK)
img(s, "pakistan-waste-contrast.jpg", Inches(6.9), Inches(1.25), Inches(5.95), Inches(5.2))
footer(s, 12)

# 13 How Pakistan
s = blank()
box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "HOW PAKISTAN ACHIEVES THIS", 12, True, GREEN)
box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.55), "Four phases — do not copy Germany’s 4 bins on day one", 22, True, INK)
phases = [
    ("01  Two-bin home", "Green = organic/wet. Blue = dry recyclables. Master two streams before four."),
    ("02  Community hubs", "Neighbourhood collection points (Lyari/Karachi-style pilots) with scheduled pickup."),
    ("03  Formalize kabbadi", "Scrap dealers into official centres: ID, fair rates, PPE, clean buy-back."),
    ("04  Local Pfand", "Small cash refund for PET at kiryana stores. Rs 5–10 trains the same muscle as €0.25."),
]
for i, (t, b) in enumerate(phases):
    x = Inches(0.4) + i * Inches(3.2)
    rrect(s, x, Inches(1.3), Inches(3.05), Inches(5.4), WHITE)
    rect(s, x, Inches(1.3), Inches(3.05), Inches(0.12), [GREEN, NAVY, YELLOW, BROWN][i])
    box(s, x + Inches(0.15), Inches(1.6), Inches(2.75), Inches(1.3), t, 18, True, INK)
    box(s, x + Inches(0.15), Inches(3.1), Inches(2.75), Inches(3.2), b, 15, False, MUTED)
footer(s, 13)

# 14 Benefits
s = blank()
box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "BENEFITS IF WE APPLY THIS", 12, True, GREEN)
box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.5), "What Pakistan gains when sorting becomes normal", 24, True, INK)
bens = [
    (GREEN, "ENVIRONMENTAL", ["Cleaner streets and parks", "Unclogged drains → less monsoon flooding", "Less open burning → cleaner air"]),
    (NAVY, "ECONOMIC", ["Jobs in formal sorting units", "Compost revenue for agriculture", "Recyclate sales instead of imports"]),
    (YELLOW, "HEALTH", ["Fewer dengue / malaria vectors", "Less toxic fume exposure", "Safer work for waste pickers with PPE"]),
    (BROWN, "ENERGY", ["Biogas from organic biomass", "Lower landfill methane", "Waste-to-energy potential"]),
]
for i, (c, t, items) in enumerate(bens):
    r, col = divmod(i, 2)
    x = Inches(0.45) + col * Inches(6.4)
    y = Inches(1.2) + r * Inches(2.85)
    card(s, x, y, Inches(6.2), Inches(2.7))
    rect(s, x, y, Inches(6.2), Inches(0.1), c)
    box(s, x + Inches(0.2), y + Inches(0.2), Inches(5.8), Inches(0.35), t, 13, True, c)
    bullets(s, items, x + Inches(0.2), y + Inches(0.65), Inches(5.8), Inches(1.85), 15)
footer(s, 14)

# 15 Consequences
s = blank()
box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "IF WE DO NOT APPLY THIS", 12, True, DANGER)
box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.55), "The punishment of not sorting is already here", 24, True, DANGER)
cons = [
    ("Monsoon flooding", "Polythene jams drains. Cities drown on modest rain because gutters are landfills."),
    ("Soil & groundwater", "Open dumps leach into wells and fields. Slow, then permanent."),
    ("Health budget", "Dengue wards, asthma, diarrhoea — the poor pay first."),
    ("Land value & tourism", "Burning heaps destroy neighbourhood prices and visitor appeal."),
]
for i, (t, b) in enumerate(cons):
    r, c = divmod(i, 2)
    x = Inches(0.45) + c * Inches(6.4)
    y = Inches(1.25) + r * Inches(2.2)
    rrect(s, x, y, Inches(6.2), Inches(2.05), RGBColor(0xFF, 0xEB, 0xEE))
    box(s, x + Inches(0.2), y + Inches(0.15), Inches(5.8), Inches(0.4), t, 16, True, DANGER)
    box(s, x + Inches(0.2), y + Inches(0.6), Inches(5.8), Inches(1.2), b, 14, False, INK)
rrect(s, Inches(0.45), Inches(5.75), Inches(12.4), Inches(1.2), NAVY)
box(s, Inches(0.65), Inches(5.95), Inches(12.0), Inches(0.85), "Germany’s ticket is €50–€2,500. Pakistan’s ticket is flooded homes, dengue, and a generation that thinks filth is normal.", 16, True, WHITE)
footer(s, 15)

# 16 PK policy
s = blank()
box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "PAKISTAN POLICY · FINES & REWARDS", 12, True, GREEN)
box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.5), "Enforcement on one side, incentives on the other", 24, True, INK)
rrect(s, Inches(0.45), Inches(1.2), Inches(6.15), Inches(5.55), RGBColor(0xB7, 0x1C, 0x1C))
box(s, Inches(0.65), Inches(1.4), Inches(5.75), Inches(0.3), "PENALTIES IF THE RULE IS APPLIED", 12, True, YELLOW)
bullets(s, [
    "Fines on commercial markets that sweep mixed waste into streets or nullahs.",
    "Household citations for dumping in public plots / drains.",
    "Repeat building/lane fees when a street keeps mixing streams.",
    "Refuse to collect unsorted commercial bins until separated (sticker method).",
], Inches(0.65), Inches(1.9), Inches(5.75), Inches(4.4), 15, WHITE)
rrect(s, Inches(6.8), Inches(1.2), Inches(6.1), Inches(5.55), GREEN)
box(s, Inches(7.0), Inches(1.4), Inches(5.7), Inches(0.3), "INCENTIVES IF A COMMUNITY APPLIES", 12, True, YELLOW)
bullets(s, [
    "Utility-bill discounts or local rebates for high-recycling neighbourhoods.",
    "Official buy-back rates at formal kabbadi centres — cash for clean PET.",
    "School sorting drills in Ring Green clubs and board curriculum.",
    "Volunteer hours counted as civic credit for youth programmes.",
], Inches(7.0), Inches(1.9), Inches(5.7), Inches(4.4), 15, WHITE)
footer(s, 16)

# 17 Youth
s = blank()
box(s, Inches(0.45), Inches(0.22), Inches(12), Inches(0.28), "ACTION CHECKLIST · MEMBERS", 12, True, GREEN)
box(s, Inches(0.45), Inches(0.5), Inches(12), Inches(0.55), "Start this week — not after a by-law", 24, True, INK)
acts = [
    ("01", "Sort at home today", "Two bags: organic scraps vs dry plastics/paper/metal. Teach the cook and the youngest sibling."),
    ("02", "Protect the drain", "No wrapper in the gali nala. A blocked drain is a shared flood."),
    ("03", "Ask for bins", "Advocate UC / municipal colour bins at the chowk and the school gate."),
    ("04", "Show up", "Cleanup drives. One kabbadi contact as the lane’s buy-back point."),
]
for i, (n, t, b) in enumerate(acts):
    r, c = divmod(i, 2)
    x = Inches(0.45) + c * Inches(6.4)
    y = Inches(1.25) + r * Inches(2.5)
    card(s, x, y, Inches(6.2), Inches(2.35))
    box(s, x + Inches(0.2), y + Inches(0.15), Inches(5.8), Inches(0.4), n + "  " + t, 16, True, NAVY)
    box(s, x + Inches(0.2), y + Inches(0.7), Inches(5.8), Inches(1.4), b, 14, False, INK)
footer(s, 17)

# 18 Close
s = prs.slides.add_slide(prs.slide_layouts[6])
img(s, "green-pakistan-future.jpg", 0, 0, W, H)
rect(s, 0, 0, Inches(8.4), H, RGBColor(0x0D, 0x1B, 0x12))
box(s, Inches(0.5), Inches(0.45), Inches(7.5), Inches(0.3), "RING GREEN PAKISTAN", 12, True, YELLOW)
box(s, Inches(0.5), Inches(1.6), Inches(7.5), Inches(1.4), "Building a sustainable future", 32, True, WHITE)
box(s, Inches(0.5), Inches(3.2), Inches(7.3), Inches(1.2), "Small daily actions lead to country-wide transformation. Let’s make Pakistan clean and green.", 18, False, WHITE)
for i, (t, d) in enumerate([("Q&A", "Bins, Pfand, and your gali pilot"), ("Connect", "@RingGreenPakistan"), ("Thank you", "4th Session · Mülltrennung & Pfand")]):
    x = Inches(0.5) + i * Inches(4.2)
    rrect(s, x, Inches(5.15), Inches(4.0), Inches(1.45), RGBColor(0x1B, 0x5E, 0x20))
    box(s, x + Inches(0.15), Inches(5.3), Inches(3.7), Inches(0.3), t.upper(), 11, True, YELLOW)
    box(s, x + Inches(0.15), Inches(5.7), Inches(3.7), Inches(0.7), d, 14, True, WHITE)
footer(s, 18)

out = Path(__file__).resolve().parent / "Ring_Green_Pakistan_Muelltrennung.pptx"
prs.save(out)
print("Wrote", out)
old = Path(__file__).resolve().parent / "Pakistan_Documented_Economy.pptx"
if old.exists():
    old.unlink()
    print("Removed", old)
